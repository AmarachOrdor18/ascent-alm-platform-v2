"""Organises the UAT input files into one folder per party and rebuilds the Group market-data workbook
(FX, curves and policy rates for NGN, GHS and KES)."""
import csv
import shutil
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

REPO = Path(r"C:\Users\AmarachiOrdor\Documents\Consulting Demo\ascent-alm-v2")
OUT = Path(r"C:\Users\AmarachiOrdor\Documents\Consulting Demo\test-files")
AS_OF = "2026-07-31"

group, ng, gh, ke = (OUT / d for d in ("1_Group_market_data", "2_Nigeria_NG", "3_Ghana_GH", "4_Kenya_KE"))
for d in (group, ng, gh, ke):
    d.mkdir(parents=True, exist_ok=True)

# Nigeria: move the files already in the root into its folder.
for name in ("loans.csv", "deposits.csv", "treasury.csv", "gl_ledger.csv", "balance-history.csv", "account-history.csv"):
    if (OUT / name).exists():
        shutil.move(str(OUT / name), ng / name)

# Ghana: the same four department files, 31 Jul 2026.
for name in ("loans.csv", "deposits.csv", "treasury.csv", "gl_ledger.csv"):
    shutil.copy(REPO / "sample-data" / "ghana" / AS_OF / name, gh / name)

# Kenya: positions per department + GL, plus two deliberately broken files for negative tests.
ke_src = REPO / "demo_data" / "affiliates" / "Ecobank Kenya"
for src, dst in (
    ("KE_positions_LOANS_2026-07-31.csv", "loans.csv"),
    ("KE_positions_DEPOSITS_2026-07-31.csv", "deposits.csv"),
    ("KE_positions_TREASURY_2026-07-31.csv", "treasury.csv"),
    ("KE_gl_trial_balance_2026-07.csv", "gl_ledger.csv"),
    ("KE_validation_failures.csv", "NEGATIVE_TEST_positions_with_errors.csv"),
    ("KE_gl_trial_balance_OUT_OF_BALANCE.csv", "NEGATIVE_TEST_gl_out_of_balance.csv"),
):
    shutil.copy(ke_src / src, ke / dst)

# Group market data: one workbook, three sheets.
policy = []
src_policy = OUT / "policy-rate.csv"
with open(src_policy, newline="") as f:
    policy = [(r["asofdate"], float(r["value"])) for r in csv.DictReader(f)]

wb = Workbook()
head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B3B4F")


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = head, fill
    for r in rows:
        ws.append(r)
    for col, w in zip("ABCDE", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


fx = wb.active
fx.title = "FX rates"
sheet(fx, ["base", "quote", "rate", "asOfDate"], [
    ["USD", "NGN", 1550.00, AS_OF], ["USD", "GHS", 15.20, AS_OF], ["USD", "KES", 129.40, AS_OF],
    ["USD", "XOF", 605.50, AS_OF], ["USD", "EUR", 0.92, AS_OF], ["USD", "ZMW", 26.10, AS_OF],
], [10, 10, 12, 14])

curves = {
    ("NGN", "NGN-NIBOR"): [(1, 18.5), (30, 19.0), (90, 19.8), (180, 20.4), (365, 21.0), (1095, 17.5), (1825, 16.0), (3650, 15.2)],
    ("GHS", "GHS-GHREF"): [(1, 26.0), (30, 26.8), (90, 27.5), (180, 28.0), (365, 28.4), (1095, 24.0), (1825, 22.5), (3650, 21.0)],
    ("KES", "KES-KESONIA"): [(1, 12.6), (30, 12.9), (90, 13.4), (180, 13.8), (365, 14.2), (1095, 14.9), (1825, 15.3), (3650, 15.8)],
    ("USD", "USD-SOFR"): [(1, 5.30), (30, 5.32), (90, 5.35), (180, 5.20), (365, 4.95), (1095, 4.30), (1825, 4.15), (3650, 4.20)],
}
yc = wb.create_sheet("Yield curves")
sheet(yc, ["currency", "curveCode", "asOfDate", "tenorDays", "ratePercent"],
      [[c, code, AS_OF, d, r] for (c, code), pts in curves.items() for d, r in pts], [10, 15, 14, 11, 12])

rows = [["CBN_MPR", "NG", d, v] for d, v in policy]
gh_rate, ke_rate = 29.0, 13.0
for d, _ in policy:
    rows.append(["BOG_POLICY", "GH", d, gh_rate])
    rows.append(["CBK_CBR", "KE", d, ke_rate])
    gh_rate, ke_rate = round(max(25.0, gh_rate - 0.15), 2), round(max(10.75, ke_rate - 0.10), 2)
pr = wb.create_sheet("Policy rates")
sheet(pr, ["indicatorCode", "countryCode", "asOfDate", "value"], rows, [16, 13, 14, 10])

wb.save(group / "market_data_2026-07-31.xlsx")
old = OUT / "market_data_2026-07-31.xlsx"
if old.exists():
    old.unlink()
shutil.move(str(src_policy), group / "policy-rate_NG_only.csv")

(OUT / "README.txt").write_text(
    "Ascent ALM - UAT input files (as of 31 Jul 2026)\n\n"
    "1_Group_market_data  Upload ONCE as Organisation Admin: Data > Sources > Upload data > Group 'Market data'.\n"
    "                     market_data_2026-07-31.xlsx = FX rates, yield curves (NGN, GHS, KES, USD), policy rates (CBN, BoG, CBK).\n"
    "2_Nigeria_NG         loans/deposits/treasury (Positions), gl_ledger (General ledger),\n"
    "                     balance-history + account-history (behavioural models, optional).\n"
    "3_Ghana_GH           loans/deposits/treasury (Positions), gl_ledger (General ledger).\n"
    "4_Kenya_KE           loans/deposits/treasury (Positions), gl_ledger (General ledger).\n"
    "                     NEGATIVE_TEST_* files are meant to FAIL (row errors; GL out of balance).\n\n"
    "Behavioural-model history is supplied for Nigeria only.\n",
    encoding="utf-8",
)
print("organised:", [p.name for p in OUT.iterdir()])
