"""Builds the Ecobank ALM clarification response in the exact format of the Nepal BRD: the Nepal file is the
template (cover page with the Qucoon logo, Biome font, Heading 1 style, justified body, navy-header tables,
sign-off table). Only the content changes."""
import copy
import sys
from docx import Document
from docx.oxml.ns import qn

SRC = sys.argv[1]
OUT = sys.argv[2]
d = Document(SRC)
body = d.element.body
paras = d.paragraphs

# ── Templates lifted from the Nepal document ──
T_HEADING = copy.deepcopy(paras[18]._p)          # "1. Introduction"
T_BODY = copy.deepcopy(paras[19]._p)             # bold lead + text
T_PLAIN = copy.deepcopy(paras[29]._p)            # body paragraph, no bold lead
T_BULLET = copy.deepcopy(paras[44]._p)           # List Bullet
T_TABLE = copy.deepcopy(d.tables[0]._tbl)        # navy header table
T_SIGN = copy.deepcopy(d.tables[4]._tbl)         # sign-off table


def set_text(p_el, text):
    """Put `text` in the paragraph's first run, drop the other runs."""
    runs = p_el.findall(qn('w:r'))
    for r in runs[1:]:
        p_el.remove(r)
    t = runs[0].find(qn('w:t'))
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return p_el


# ── Cover: same layout, new words ──
def cover(i, text):
    p = paras[i]._p
    set_text(p, text)


cover(9, 'Response to Clarification Questions')
cover(10, 'Ecobank Group – Asset & Liability Management (ALM) Platform')
cover(14, '19th October 2026')

# ── Drop the Nepal body (everything after the cover page break), keep the section settings ──
cover_end = paras[17]._p
el = cover_end.getnext()
while el is not None and el.tag != qn('w:sectPr'):
    nxt = el.getnext()
    body.remove(el)
    el = nxt
sect = body.find(qn('w:sectPr'))


def add(el):
    sect.addprevious(el)


def heading(text):
    add(set_text(copy.deepcopy(T_HEADING), text))


def lead(bold, text):
    p = copy.deepcopy(T_BODY)
    r1, r2 = p.findall(qn('w:r'))[:2]
    r1.find(qn('w:t')).text = bold + ' '
    r2.find(qn('w:t')).text = text
    for r in p.findall(qn('w:r'))[2:]:
        p.remove(r)
    add(p)


def plain(text):
    p = set_text(copy.deepcopy(T_PLAIN), text)
    # No first-line indent - flush with the other paragraphs.
    ppr = p.find(qn('w:pPr'))
    if ppr is not None:
        for ind in ppr.findall(qn('w:ind')):
            ppr.remove(ind)
        ppr.append(ppr.makeelement(qn('w:ind'), {qn('w:left'): '0', qn('w:firstLine'): '0'}))
    add(p)


def bullet(text):
    add(set_text(copy.deepcopy(T_BULLET), text))


def spacer():
    p = copy.deepcopy(T_PLAIN)
    for r in p.findall(qn('w:r')):
        p.remove(r)
    add(p)


def table(header, rows, widths):
    """A table in the Nepal style: navy header row, bordered 9pt body rows, given column widths (twips)."""
    tbl = copy.deepcopy(T_TABLE)
    trs = tbl.findall(qn('w:tr'))
    head_tr, body_tr = trs[0], trs[1]
    for tr in trs:
        tbl.remove(tr)
    grid = tbl.find(qn('w:tblGrid'))
    for g in list(grid):
        grid.remove(g)
    total = sum(widths)
    tbl.find(qn('w:tblPr')).find(qn('w:tblW')).set(qn('w:w'), str(total))
    for w in widths:
        g = copy.deepcopy(trs[0].find(qn('w:tc')))  # placeholder to get namespace right
        col = grid.makeelement(qn('w:gridCol'), {qn('w:w'): str(w)})
        grid.append(col)

    def make_row(template_tr, values):
        tr = copy.deepcopy(template_tr)
        cells = tr.findall(qn('w:tc'))
        proto = cells[0]
        for c in cells:
            tr.remove(c)
        for value, w in zip(values, widths):
            tc = copy.deepcopy(proto)
            tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(w))
            p = tc.find(qn('w:p'))
            set_text(p, value)
            tr.append(tc)
        return tr

    tbl.append(make_row(head_tr, header))
    for r in rows:
        tbl.append(make_row(body_tr, r))
    add(tbl)


# ══ Content ══
heading('1. Introduction')
lead('Purpose.', 'This document is Qucoon’s response to the clarification questions raised by Ecobank Group on the '
     'behavioural modelling and funds transfer pricing (FTP) capabilities of the proposed Asset & Liability Management '
     '(ALM) platform, Ascent ALM. Each answer states what the platform does and where Ecobank can see it working.')
lead('Background.', 'Ecobank Group asked whether the proposed solution derives behavioural assumptions natively, applies '
     'them across all ALM calculations without re-entry, and provides a complete FTP and profitability engine on the '
     'same cashflows as liquidity and interest-rate risk. The answers below are confirmed against the working platform.')
lead('How to read the answers.', 'Every response gives a direct answer, a short explanation and the screen in the platform '
     'where the capability can be tested end to end by Ecobank’s own users, without Qucoon support.')

heading('2. Platform Access')
lead('Link.', 'The capabilities in this document can be validated directly on the Ascent ALM user-acceptance environment at '
     'https://ascent-alm-v2.vercel.app (any modern browser; no installation).')
lead('Sign-in.', 'Use the accounts below, each set up with a different role so the business team can see maker-checker '
     'controls working. The password is provided separately with this document. Dedicated named accounts for Ecobank’s '
     'reviewers can be created on request, and are recommended for any extended review.')
table(
    ['Role', 'Sign-in (email)', 'Use it to validate'],
    [
        ['Organisation Admin', 'adaeze.okonkwo@ecobank.com', 'Group set-up, market data, Group assumptions, users'],
        ['ALM Analyst', 'chinwe.okafor@ecobank.com', 'Behavioural models (B1–B5), runs, FTP and profitability results (F1–F4)'],
        ['Approver', 'kofi.mensah@ecobank.com', 'Approving model-derived assumptions and run sign-offs (B6)'],
        ['Affiliate Admin (Nigeria)', 'ngadmin.test@ecobank.com', 'An affiliate’s own view: data, its own assumptions, readiness'],
        ['Viewer', 'yaw.boateng@ecobank.com', 'Read-only access to results and reports'],
    ],
    [2300, 3000, 3634],
)
lead('Guide.', 'A step-by-step validation walkthrough (ALM_UAT_Walkthrough) accompanies this document; the numbered '
     'sections referenced below can be followed without Qucoon support. In the platform, choose the affiliate with Scope '
     'and the month-end with As of at the top of the page.')
spacer()

heading('3. Summary of Responses')
table(
    ['ID', 'Question', 'Response'],
    [
        ['B1', 'Native behavioural modelling engine deriving assumptions from historical data', 'Yes'],
        ['B2', 'Regression, time-series and other techniques run internally - no Python, R, Excel or third-party tools', 'Yes'],
        ['B3', 'Core and non-core deposits, stability, decay rates and runoff profiles from historical data', 'Yes'],
        ['B4', 'Loan prepayment, early withdrawal, deposit renewal and repricing betas', 'Yes'],
        ['B5', 'Calibration, back-testing and validation entirely within the platform', 'Yes'],
        ['B6', 'Approved assumptions applied automatically to liquidity, IRRBB, FTP, stress testing and forecasting', 'Yes'],
        ['B7', 'Behavioural modelling included in the licence and implementation scope', 'Fully included'],
        ['F1', 'Complete FTP engine - matched maturity, pooled funding, liquidity premiums, basis and currency adjustments', 'Yes'],
        ['F2', 'FTP at transaction level, aggregated by product, business unit, affiliate and Group', 'Yes'],
        ['F3', 'FTP uses the same contractual and behavioural cashflows as liquidity and IRRBB', 'Yes'],
        ['F4', 'Allocation of funding costs, operating expenses, expected credit losses and capital charges', 'Yes'],
    ],
    [700, 6634, 1600],
)
spacer()

heading('4. Behavioural Modelling & Assumptions')
table(
    ['ID', 'Question', 'Response', 'Where to see it in the platform'],
    [
        ['B1', 'Does the proposed solution include a native behavioural modelling engine capable of analysing historical '
               'transactional and balance data and independently deriving behavioural assumptions?',
         'Yes. The platform includes a native behavioural modelling engine that ingests historical transactional and '
         'balance data, performs statistical analysis, and derives behavioural assumptions independently within the platform.',
         'Data → Upload data (balance and account history), then Assumptions → Behavioural models → Calibrate.'],
        ['B2', 'Can the system perform statistical regression, time-series analysis and other relevant modelling techniques '
               'internally, without requiring Python, R, Excel or third-party analytical software?',
         'Yes. All statistical modelling including regression analysis, time-series decomposition and curve fitting is '
         'performed natively within the platform. No external tools such as Python, R, Excel or third-party software are required.',
         'Behavioural models → Fit model: R², RMSE, MAPE, p-values and 95% confidence intervals are shown on screen.'],
        ['B3', 'Can the solution calculate core and non-core deposit balances, deposit stability, decay rates and expected '
               'runoff profiles using historical data?',
         'Yes. The engine segments deposits into core and non-core categories based on historical behaviour, calculates '
         'deposit stability metrics, derives decay rates and generates expected runoff profiles. Segmentation is '
         'configurable by deposit type, client type, tenor, currency and affiliate.',
         'Model “Deposit stability & run-off”: core method (minimum balance or volatility), segment by product, '
         'customer type, currency or tenor band; decay rate, half-life and runoff chart.'],
        ['B4', 'Can the solution independently estimate loan prepayment rates, early withdrawal behaviour, deposit renewal '
               'rates and repricing betas?',
         'Yes. The behavioural engine models loan prepayment rates, early withdrawal patterns, deposit renewal behaviour and '
         'repricing betas using historical data. Outputs are model-derived assumptions that can be reviewed, adjusted and '
         'approved by authorised users before application.',
         'Models “Loan prepayment” (CPR), “Term deposit” (renewal and early withdrawal) and '
         '“Repricing beta” (beta, lag, fit statistics against the central bank policy rate).'],
        ['B5', 'Can users calibrate, back-test and validate behavioural models entirely within the proposed solution?',
         'Yes. The platform provides built-in calibration, back-testing and validation tools. Users compare model outputs '
         'against observed outcomes, assess model fit, review confidence intervals and adjust parameters within the platform. '
         'Model versions are tracked with full audit trails.',
         'Back-test hold-out and tolerance (MAE, RMSE, MAPE, bias, Pass/Fail); Behavioural models → Model register; '
         'Admin → Audit log.'],
        ['B6', 'Once approved, are behavioural assumptions automatically applied to liquidity risk, IRRBB, FTP, stress testing '
               'and forecasting without manual transfer or re-entry?',
         'Yes. Once behavioural assumptions are approved through the platform’s maker-checker workflow, they are '
         'automatically applied across all dependent modules - liquidity risk, IRRBB, FTP, stress testing and forecasting. '
         'No manual transfer, re-entry or file export is required.',
         'Model register → Propose as assumption → Approvals (a second user approves) → Run now; '
         'Run & results → Behavioural analysis lists the approved version each run applied.'],
        ['B7', 'Is the complete behavioural modelling, calibration and back-testing functionality included in the proposed '
               'licence and implementation scope, or does it require an additional module, external consultancy or a '
               'separate modelling platform?',
         'Fully included. The behavioural modelling engine, calibration tools and back-testing functionality are all part of '
         'the core platform scope and the proposed commercial terms.',
         'All of the above runs in the same platform - no separate module, export or external tool.'],
    ],
    [600, 2650, 3384, 2300],
)
spacer()

heading('5. FTP & Profitability')
table(
    ['ID', 'Question', 'Response', 'Where to see it in the platform'],
    [
        ['F1', 'Is the complete FTP engine included in the proposed commercial solution, including matched-maturity FTP, '
               'pooled funding, liquidity premiums and basis/currency adjustments?',
         'Yes. The FTP engine is fully included in the proposed scope and commercial terms. It supports matched-maturity FTP, '
         'pooled funding rate allocation, liquidity premium assignment, and basis and currency adjustments.',
         'Assumptions → Rules & scenarios → FTP method & curve and FTP add-ons; FTP & profitability → '
         'Transfer pricing → Compare methods.'],
        ['F2', 'Can the solution calculate FTP at individual transaction level and automatically aggregate the results by '
               'product, business unit, affiliate and Group?',
         'Yes. FTP is calculated at the individual instrument/transaction level and automatically aggregated by product, '
         'business unit, affiliate and Group. Users can drill down from Group-level profitability to the underlying '
         'transaction detail.',
         'Transfer pricing drill-down: Group → affiliate → business unit → product → positions; each '
         'line shows base rate, add-ons and all-in rate.'],
        ['F3', 'Does the FTP engine use the same contractual and behavioural cashflows as the liquidity and IRRBB engines?',
         'Yes. All engines operate on a single, consistent set of cashflow data. The contractual and behavioural cashflows '
         'used by the FTP engine are the same cashflows used by the liquidity and IRRBB engines. There is no duplication or '
         'separate data pipeline.',
         'Approve a prepayment assumption and re-run: the behavioural gap, loan FTP rates and EVE all move together.'],
        ['F4', 'Can the system allocate funding costs, operating expenses, expected credit losses and capital charges to '
               'support product and customer profitability analysis?',
         'Yes. The profitability module supports multi-layer cost allocation including funding costs (via FTP), operating '
         'expense allocation, expected credit loss attribution and capital charges. This enables full product-level and '
         'customer-level profitability analysis.',
         'FTP & profitability → Profitability: waterfall to risk-adjusted profit and RAROC, Product / Customer view, '
         'cost pool reconciliation.'],
    ],
    [600, 2650, 3384, 2300],
)
spacer()

heading('6. Demonstration Environment')
lead('Set-up.', 'The answers above can be tested on the Ascent ALM user-acceptance environment, configured for Ecobank Group '
     'with its affiliates onboarded through the platform’s own onboarding flow, Group market data loaded once for all '
     'affiliates, and Group-default assumptions approved through maker-checker.')
table(
    ['Affiliate', 'Status', 'Month-end runs', 'Notes'],
    [
        ['Ecobank Nigeria', 'Live', '2', '31 Jul 2026 (signed off) and 31 Aug 2026; behavioural, FTP and profitability results'],
        ['Ecobank Ghana', 'Live', '1', '31 Jul 2026; onboarded through the same affiliate flow'],
        ['Ecobank Group', 'Consolidated', '-', 'No separate Group run: the Group view consolidates each Live affiliate’s '
         'official run for the month-end (e.g. 31 Jul 2026 = Nigeria + Ghana)'],
    ],
    [2000, 1400, 900, 4634],
)
spacer()

heading('7. Notes and Assumptions')
lead('Data loading.', 'In the test environment historical balance, account and position data are loaded through the '
     'platform’s data sources by file upload, checked and admitted by Ecobank users. Automated connectors to source '
     'systems are configured during implementation.')
lead('Governance.', 'Every assumption change, model proposal and run sign-off passes through maker-checker approval; the '
     'person who raises a change cannot approve it, and every action is recorded in an append-only audit log.')

heading('8. Next Steps')
lead('Next steps.', 'Ecobank’s business team validates the capabilities above on the user-acceptance environment (Section 2) '
     'using the walkthrough provided. Qucoon is available to clarify any point or set up named reviewer accounts.')
plain('Please let us know if any further clarification is required.')

heading('9. Sign-off')
sign = copy.deepcopy(T_SIGN)
trs = sign.findall(qn('w:tr'))
names = [
    ('', 'Head of ALM', 'Ecobank Group'),
    ('', 'Group Chief Risk Officer', 'Ecobank Group'),
    ('', 'Engagement Lead', 'Qucoon Limited'),
]
for tr in trs[1:]:
    sign.remove(tr)
for name, role, org in names:
    tr = copy.deepcopy(trs[1])
    cells = tr.findall(qn('w:tc'))
    for tc, val in zip(cells, [name, role, org, '', '']):
        p = tc.find(qn('w:p'))
        runs = p.findall(qn('w:r'))
        if runs:
            set_text(p, val)
    sign.append(tr)
add(sign)

d.core_properties.title = 'Response to Clarification Questions - Ecobank Group ALM Platform'
d.save(OUT)
print('->', OUT)
