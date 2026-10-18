"""Builds test-files/1_Group_market_data/holiday_calendars_2026-2027.xlsx (and a .csv copy): public holidays for
Nigeria, Ghana and Kenya for 2026 and 2027, uploaded on Data > Market data > Holiday calendars > Upload calendar.
Islamic holidays follow the moon - their dates are estimates, marked '(est.)', and are confirmed by government
notice each year."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(r"C:\Users\AmarachiOrdor\Documents\Consulting Demo\test-files\1_Group_market_data")
WEEKEND = "Sat,Sun"

CALENDARS = {
    ("NG", "Nigeria (CBN) business days"): [
        ("2026-01-01", "New Year's Day"), ("2026-03-20", "Eid al-Fitr (est.)"), ("2026-03-21", "Eid al-Fitr holiday (est.)"),
        ("2026-04-03", "Good Friday"), ("2026-04-06", "Easter Monday"), ("2026-05-01", "Workers' Day"),
        ("2026-05-27", "Eid al-Adha (est.)"), ("2026-05-28", "Eid al-Adha holiday (est.)"), ("2026-06-12", "Democracy Day"),
        ("2026-08-26", "Eid al-Mawlid (est.)"), ("2026-10-01", "Independence Day"), ("2026-12-25", "Christmas Day"),
        ("2026-12-28", "Boxing Day (observed)"),
        ("2027-01-01", "New Year's Day"), ("2027-03-10", "Eid al-Fitr (est.)"), ("2027-03-26", "Good Friday"),
        ("2027-03-29", "Easter Monday"), ("2027-05-03", "Workers' Day (observed)"), ("2027-05-17", "Eid al-Adha (est.)"),
        ("2027-06-14", "Democracy Day (observed)"), ("2027-08-16", "Eid al-Mawlid (est.)"), ("2027-10-01", "Independence Day"),
        ("2027-12-27", "Christmas Day (observed)"), ("2027-12-28", "Boxing Day (observed)"),
    ],
    ("GH", "Ghana (BoG) business days"): [
        ("2026-01-01", "New Year's Day"), ("2026-01-07", "Constitution Day"), ("2026-03-06", "Independence Day"),
        ("2026-03-20", "Eid al-Fitr (est.)"), ("2026-04-03", "Good Friday"), ("2026-04-06", "Easter Monday"),
        ("2026-05-01", "May Day"), ("2026-05-27", "Eid al-Adha (est.)"), ("2026-07-01", "Republic Day"),
        ("2026-08-04", "Founders' Day"), ("2026-09-21", "Kwame Nkrumah Memorial Day"), ("2026-12-04", "Farmers' Day"),
        ("2026-12-25", "Christmas Day"), ("2026-12-28", "Boxing Day (observed)"),
        ("2027-01-01", "New Year's Day"), ("2027-01-07", "Constitution Day"), ("2027-03-08", "Independence Day (observed)"),
        ("2027-03-10", "Eid al-Fitr (est.)"), ("2027-03-26", "Good Friday"), ("2027-03-29", "Easter Monday"),
        ("2027-05-03", "May Day (observed)"), ("2027-05-17", "Eid al-Adha (est.)"), ("2027-07-01", "Republic Day"),
        ("2027-08-04", "Founders' Day"), ("2027-09-21", "Kwame Nkrumah Memorial Day"), ("2027-12-03", "Farmers' Day"),
        ("2027-12-27", "Christmas Day (observed)"), ("2027-12-28", "Boxing Day (observed)"),
    ],
    ("KE", "Kenya (CBK) business days"): [
        ("2026-01-01", "New Year's Day"), ("2026-03-20", "Eid al-Fitr (est.)"), ("2026-04-03", "Good Friday"),
        ("2026-04-06", "Easter Monday"), ("2026-05-01", "Labour Day"), ("2026-06-01", "Madaraka Day"),
        ("2026-10-10", "Mazingira Day"), ("2026-10-20", "Mashujaa Day"), ("2026-12-12", "Jamhuri Day"),
        ("2026-12-25", "Christmas Day"), ("2026-12-26", "Boxing Day"),
        ("2027-01-01", "New Year's Day"), ("2027-03-10", "Eid al-Fitr (est.)"), ("2027-03-26", "Good Friday"),
        ("2027-03-29", "Easter Monday"), ("2027-05-01", "Labour Day"), ("2027-06-01", "Madaraka Day"),
        ("2027-10-11", "Mazingira Day (observed)"), ("2027-10-20", "Mashujaa Day"), ("2027-12-13", "Jamhuri Day (observed)"),
        ("2027-12-25", "Christmas Day"), ("2027-12-27", "Boxing Day (observed)"),
    ],
}

header = ["countryCode", "calendarName", "date", "holidayName", "weekendDays"]
rows = [[cc, name, d, h, WEEKEND] for (cc, name), days in CALENDARS.items() for d, h in days]

wb = Workbook()
ws = wb.active
ws.title = "Holidays"
ws.append(header)
for c in ws[1]:
    c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B3B4F")
for r in rows:
    ws.append(r)
for col, w in zip("ABCDE", [13, 30, 12, 32, 13]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
wb.save(OUT / "holiday_calendars_2026-2027.xlsx")

with open(OUT / "holiday_calendars_2026-2027.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(header)
    w.writerows(rows)

readme = OUT.parent / "README.txt"
text = readme.read_text(encoding="utf-8")
if "holiday_calendars_2026-2027" not in text:
    text += (
        "\nHoliday calendars:\n"
        "  1_Group_market_data/holiday_calendars_2026-2027.xlsx (or .csv)   public holidays for NG, GH and KE, 2026-2027.\n"
        "  Upload on Data > Market data > Holiday calendars > Upload calendar. Eid dates are estimates (marked est.).\n"
    )
    readme.write_text(text, encoding="utf-8")
print(len(rows), "holidays written")
