"""Builds test-files/1_Group_market_data/Group_assumptions_setup_values.xlsx - the values to type into
the organisation onboarding "Assumptions" step (one sheet per assumption). Product classes / codes are the
ones used in the Nigeria/Ghana/Kenya test position files (COA-11 cash, COA-12 interbank, COA-13 government
securities, COA-14 loans, COA-22 deposits, COA-23 borrowings, COA-31 capital)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(r"C:\Users\AmarachiOrdor\Documents\Consulting Demo\test-files\1_Group_market_data\Group_assumptions_setup_values.xlsx")
HEAD, FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B3B4F")
WRAP = Alignment(wrap_text=True, vertical="top")

SHEETS = {
    "Read me": (["#", "Assumption", "Needed?", "Approval?", "Where"], [
        [1, "Transfer pricing method & curve", "Optional", "No - saved = in use", "Onboarding Step 4 > Set up (or Assumptions > Rules & scenarios > FTP method & curve)"],
        [2, "Transfer pricing add-ons", "Optional", "No - saved = in use", "Step 4 > Set up"],
        [3, "Operating cost allocation", "Needed for profitability", "Yes - an Approver approves", "Step 4 > Set up"],
        [4, "Expected credit loss", "Needed for profitability", "Yes - an Approver approves", "Step 4 > Set up"],
        [5, "Capital charge", "Needed for profitability", "Yes - an Approver approves", "Step 4 > Set up"],
        [6, "Product characteristics", "Optional (files already carry liquidity treatment)", "Yes - an Approver approves", "Step 4 > Set up"],
        [7, "Deposit behaviour", "Optional - a standard pattern already ships", "Yes, if changed", "Leave as shipped; replace later with a fitted model"],
        ["", "Approvals", "", "", "Sign in as kofi.mensah@ecobank.com (Approver) > Approvals > Waiting for me > Approve each request"],
    ]),
    "1 FTP method": (["Field", "Value", "Note"], [
        ["Rule name", "Group FTP - Matched maturity", ""],
        ["Default method", "Matched maturity", "Applies to anything not listed below"],
        ["Assignment: product class COA-14 (Loans)", "Matched maturity · Curve: - affiliate default -", "Each loan funded at its own repayment dates"],
        ["Assignment: product class COA-22 (Deposits)", "Matched maturity · Curve: - affiliate default -", "Deposits credited on their behavioural tenor"],
        ["Assignment: product class COA-13 (Government securities)", "Repricing term · Curve: - affiliate default -", ""],
        ["Assignment: product class COA-12 (Interbank)", "Repricing term · Curve: - affiliate default -", ""],
        ["Funding pools (optional)", "Leave empty", "Only needed to test the Pooled method"],
    ]),
    "2 FTP add-ons": (["Type", "Applies to", "Method", "Value", "Note"], [
        ["Rule name: Group FTP add-ons", "", "", "", ""],
        ["Liquidity premium", "COA-14 (Loans)", "Fixed", "25 bps", "Charged on loans"],
        ["Liquidity premium", "COA-22 (Deposits)", "Fixed", "15 bps", "Credited to stable funding"],
        ["Currency basis", "Currency USD", "Fixed", "40 bps", "Only hits non-local-currency positions"],
    ]),
    "3 Cost allocation": (["Pool name", "Annual amount", "Driver", "Targets", "Note"], [
        ["Rule name: Group operating costs  ·  Currency: USD", "", "", "", "Converted at run-date FX for each affiliate's run"],
        ["Staff costs", "12,000,000", "Balance", "All", ""],
        ["Technology", "4,000,000", "Account count", "All", ""],
        ["Premises & branches", "6,000,000", "Balance", "All", ""],
        ["Headcount", "Leave empty", "", "", "Only needed for the Headcount driver"],
    ]),
    "4 Credit loss": (["Field / product", "PD %", "LGD %", "Note"], [
        ["Rule name: Group expected credit loss  ·  Method: PD × LGD", "", "", ""],
        ["Default (any product not listed)", "2.0", "45", ""],
        ["LOAN-RET (retail loans)", "3.0", "40", ""],
        ["LOAN-COR (corporate loans)", "1.5", "45", ""],
        ["LOAN-SME (SME loans)", "4.0", "50", ""],
        ["Use booked provisions for non-performing loans", "On", "", ""],
    ]),
    "5 Capital": (["Field / product class", "Value", "Note"], [
        ["Rule name", "Group capital charge", ""],
        ["Capital ratio", "15%", ""],
        ["Cost of capital", "18%", ""],
        ["Tax rate", "30%", "Used in RAROC"],
        ["Default risk weight", "100%", ""],
        ["COA-11 Cash & central bank", "0%", ""],
        ["COA-12 Interbank placements", "20%", ""],
        ["COA-13 Government securities", "0%", ""],
        ["COA-14 Loans", "100%", ""],
    ]),
    "6 Product characteristics": (["Product", "Currency", "LCR role", "LCR rate %", "ASF %", "RSF %", "HQLA level", "Haircut %"], [
        ["CASH", "NGN", "HQLA", "", "", "0", "Level 1", "0"],
        ["TBILL", "NGN", "HQLA", "", "", "5", "Level 1", "0"],
        ["FGNBOND", "NGN", "HQLA", "", "", "5", "Level 1", "0"],
        ["INTERBANK", "NGN", "Inflow", "50", "", "15", "", ""],
        ["DDA", "NGN", "Outflow", "10", "90", "", "", ""],
        ["TD", "NGN", "Outflow", "5", "95", "", "", ""],
        ["LOAN-RET", "NGN", "Inflow", "50", "", "85", "", ""],
        ["LOAN-COR", "NGN", "Inflow", "50", "", "85", "", ""],
        ["LOAN-SME", "NGN", "Inflow", "50", "", "85", "", ""],
        ["(then)", "", "", "", "", "", "", "Use 'Copy across currencies' to repeat for GHS and KES"],
    ]),
}

wb = Workbook()
wb.remove(wb.active)
for name, (header, rows) in SHEETS.items():
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    for i, col in enumerate("ABCDEFGH"):
        ws.column_dimensions[col].width = 46 if i == 0 else (34 if len(header) <= 5 else 14)
    ws.freeze_panes = "A2"
wb.save(OUT)

readme = OUT.parent.parent / "README.txt"
text = readme.read_text(encoding="utf-8")
if "Group_assumptions_setup_values" not in text:
    text = text.replace(
        "1_Group_market_data  Upload ONCE",
        "1_Group_market_data  Group_assumptions_setup_values.xlsx = values to type in organisation onboarding Step 4 (Assumptions).\n"
        "                     Upload ONCE",
    )
    readme.write_text(text, encoding="utf-8")
print("->", OUT)
